# 只读检查已发表 GSE282122 补充表和代码，判断哪些参考构建信息可用于整块黏膜。
# 在本机 bundled Python 运行；读 sources 原件，写同目录的结构摘要和证据索引，不执行作者代码。
from pathlib import Path
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import openpyxl

# Step1. 读入原始补充表，保留表名与实际列名的对应关系
out = Path(r"D:\IBD_ecotype\IBD_ecotype_analysis\CIBERSORTx_signature_audit\published_GSE282122_reference_20261007")
source = out / "sources"
book = openpyxl.load_workbook(source / "Table_S2_mmc3.xlsx", read_only=True, data_only=True)
print("Python:", sys.executable)
print("Sheets:", book.sheetnames)
summary = []
profiles = []
counts = []

count_sheet = book["TableS2A"]
compartment = ""
for row_id, row in enumerate(count_sheet.iter_rows(values_only=True), start=1):
    if row_id < 4:
        continue
    if row[0]:
        compartment = str(row[0])
    if isinstance(row[2], (int, float)):
        counts.append({"sheet": "TableS2A", "excel_row": row_id, "compartment_label": compartment, "cell_type_label": str(row[1]).strip(), "cells": row[2]})
pd.DataFrame(counts).to_csv(out / "published_cell_counts.csv", index=False, encoding="utf-8-sig")

# Step2. 两张表达表逐一读取；这是核对作者实际交付物，不重新构建 signature
for sheet_name in ["TableS2B", "TableS2C"]:
    sheet = book[sheet_name]
    rows = sheet.iter_rows(min_row=3, values_only=True)
    header = next(rows)
    records = list(rows)
    expression = pd.DataFrame(records, columns=header).set_index("GeneSymbol")
    values = expression.to_numpy(dtype=float)
    type_labels = [re.sub(r"\.\d+$", "", str(name)) for name in expression.columns]
    type_sizes = pd.Series(type_labels).value_counts(sort=False)
    print(sheet_name, "genes x expression columns:", expression.shape, "types:", len(type_sizes))
    print("Cell types:", type_sizes.to_dict())
    item = {"sheet": sheet_name, "sheet_title": sheet.cell(1, 1).value, "genes": len(expression), "expression_columns": expression.shape[1], "type_groups": len(type_sizes), "columns_per_type": type_sizes.to_dict(), "gene_identifier_examples": expression.index[:5].tolist(), "genes_unique": bool(expression.index.is_unique), "all_finite": bool(np.isfinite(values).all()), "min_expression": float(np.nanmin(values)), "max_expression": float(np.nanmax(values)), "column_sum_min": float(values.sum(axis=0).min()), "column_sum_max": float(values.sum(axis=0).max()), "zero_fraction": float((values == 0).mean()), "all_zero_genes": int((values == 0).all(axis=1).sum())}
    summary.append(item)
    for cell_type in type_sizes.index:
        positions = [i for i, label in enumerate(type_labels) if label == cell_type]
        subset = values[:, positions]
        correlation = np.corrcoef(subset.T)
        pairwise = correlation[np.triu_indices(len(positions), k=1)]
        profiles.append({"sheet": sheet_name, "type_label": cell_type, "expression_columns": len(positions), "genes_nonzero_in_any_column": int((subset > 0).any(axis=1).sum()), "within_type_column_correlation_min": float(pairwise.min()), "column_names": "|".join(str(expression.columns[i]) for i in positions)})
    expression.index.to_series().to_csv(out / f"{sheet_name}_gene_ids.csv", index=False, header=["gene_id"], encoding="utf-8-sig")
book.close()
pd.DataFrame(profiles).to_csv(out / "published_reference_type_columns.csv", index=False, encoding="utf-8-sig")

# Step3. 读取所有公开 notebook 的源代码，定位构建或验证信息；不运行其分析
code_index = []
for path in sorted(source.glob("*.ipynb")):
    notebook = json.loads(path.read_text(encoding="utf-8"))
    cells = []
    for i, cell in enumerate(notebook["cells"]):
        cell_source = "".join(cell.get("source", []))
        cells.append(f"\n### Cell {i} ({cell['cell_type']})\n{cell_source}\n")
        for term in ["pseudobulk", "phenoclass", "CPM", "holdout", "ground_truth", "cibresults_Epithelium", "cibresults_Lamina_propria"]:
            if term.lower() in cell_source.lower():
                code_index.append({"notebook": path.name, "cell_index": i, "cell_type": cell["cell_type"], "matched_term": term})
    (out / f"{path.stem}_readonly.txt").write_text("".join(cells), encoding="utf-8")
pd.DataFrame(code_index).to_csv(out / "notebook_evidence_index.csv", index=False, encoding="utf-8-sig")

# Step4. 保存论文方法段与下载原件校验值，便于核对证据来源
root = ET.parse(source / "PMC13573676.xml").getroot()
methods = []
for section in root.iter("sec"):
    title = section.find("title")
    if title is not None and "deconvolution" in "".join(title.itertext()).lower():
        methods.append("".join(section.itertext()))
(out / "paper_deconvolution_methods.txt").write_text("\n\n".join(methods), encoding="utf-8")
integrity = []
for path in sorted(source.iterdir()):
    if path.is_file():
        integrity.append({"file": path.name, "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
pd.DataFrame(integrity).to_csv(out / "downloaded_source_integrity.csv", index=False, encoding="utf-8-sig")
(out / "supplement_structure_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print("Saved:", out)
print("Expression column sums:", [(item["sheet"], item["column_sum_min"], item["column_sum_max"]) for item in summary])
